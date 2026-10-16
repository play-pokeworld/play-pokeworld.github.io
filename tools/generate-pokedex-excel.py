#!/usr/bin/env python3
"""
Generates docs/POKEDEX_OBTENTION.xlsx — Complete obtainment guide for all 386 Pokemon
across Kanto, Johto, and Hoenn.
"""

import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load data
with open('/tmp/pokedex_data.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

pokemon_list = data['pokemon']
locations_list = data['locations']

wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

# Color palettes & styles
FONT_NAME = 'Segoe UI'
HEADER_FILL = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color='FFFFFF')
SUBHEADER_FILL = PatternFill(start_color='2C5282', end_color='2C5282', fill_type='solid')

ZEBRA_EVEN = PatternFill(start_color='F7FAFC', end_color='F7FAFC', fill_type='solid')
ZEBRA_ODD = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

BORDER_THIN = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

TYPE_COLORS = {
    'Normal': 'A0A29F', 'Fire': 'FBA64C', 'Water': '539DDF', 'Grass': '60BE58',
    'Electric': 'F2D94E', 'Ice': '76D1C1', 'Fighting': 'CE3E5A', 'Poison': 'BA68D2',
    'Ground': 'DA7C4D', 'Flying': 'A1BBEC', 'Psychic': 'FA8582', 'Bug': '92BD2D',
    'Rock': 'C9BC8A', 'Ghost': '5F6DBC', 'Dragon': '0C6AC8', 'Dark': '595761',
    'Steel': '5795A3', 'Fairy': 'EF90E6'
}

# -----------------------------------------------------------------------------
# SHEET 1: Pokedex Complet (All 386 Pokemon)
# -----------------------------------------------------------------------------
ws1 = wb.create_sheet(title='Pokédex Complet')
ws1.views.sheetView[0].showGridLines = True

headers1 = [
    'N° Dex', 'Nom (FR)', 'Nom (EN)', 'Type 1', 'Type 2', 'Région',
    'Méthode Principale', 'Lieux Sauvages (Zones & Niveaux)',
    'Évolution (Condition / Niveau / Pierre)', 'Quêtes Associées',
    'Marché Noir (Région & Prix)', 'Pension / Bébés', 'Fossile Associé', 'BST'
]

ws1.append(headers1)
for col_num, _ in enumerate(headers1, 1):
    cell = ws1.cell(row=1, column=col_num)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws1.row_dimensions[1].height = 28

for p in pokemon_list:
    dex_str = f"#{p['dex']:03d}"
    t1 = p['type1']
    t2 = p['type2'] or ''
    reg = p['region']
    method = p['primaryMethod']

    # Wild summary
    if p['wilds']:
        loc_summaries = []
        for w in p['wilds']:
            rarity_fr = 'commun' if w['rarity'] == 'common' else ('rare' if w['rarity'] == 'rare' else 'peu commun')
            loc_summaries.append(f"{w['locName']} (Nv. {w['minLvl']}-{w['maxLvl']}, {rarity_fr})")
        wild_str = " ; ".join(loc_summaries)
    else:
        wild_str = "Introuvable à l'état sauvage"

    # Evolution summary
    if p['evolution']:
        evo_str = p['evolution']['details']
    else:
        evo_str = "Forme de base ou pas d'évolution directe"

    # Quests summary
    if p['quests']:
        q_summaries = [f"{q['title']} ({q['cat']}, Nv.{q['level']}{', Shiny!' if q['shiny'] else ''})" for q in p['quests']]
        quest_str = " ; ".join(q_summaries)
    else:
        quest_str = "-"

    # Market summary
    if p['market']:
        market_str = f"{p['market']['region'].capitalize()} ({p['market']['price']:,} ₽)".replace(',', ' ')
    else:
        market_str = "-"

    # Breeding summary
    if p['breeding']:
        breed_str = f"Reproduction à la pension avec {p['breeding']['parents']}"
    else:
        breed_str = "-"

    # Fossil summary
    if p['fossil']:
        fossil_str = f"Réanimation du {p['fossil']['name']}"
    else:
        fossil_str = "-"

    bst = sum(p['stats'].values())

    row_data = [
        dex_str, p['nameFr'], p['nameEn'], t1, t2, reg, method,
        wild_str, evo_str, quest_str, market_str, breed_str, fossil_str, bst
    ]
    ws1.append(row_data)

    current_row = ws1.max_row
    ws1.row_dimensions[current_row].height = 22
    fill_to_use = ZEBRA_EVEN if current_row % 2 == 0 else ZEBRA_ODD

    for c_idx in range(1, len(row_data) + 1):
        c = ws1.cell(row=current_row, column=c_idx)
        c.fill = fill_to_use
        c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER_THIN

        if c_idx in [1, 4, 5, 6, 7, 14]:
            c.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx in [2, 3]:
            c.font = Font(name=FONT_NAME, size=10, bold=True)
            c.alignment = Alignment(horizontal='left', vertical='center')
        else:
            c.alignment = Alignment(horizontal='left', vertical='center')

ws1.freeze_panes = 'A2'
ws1.auto_filter.ref = ws1.dimensions

# -----------------------------------------------------------------------------
# SHEET 2: Lieux Sauvages (Locations & Wild Encounters)
# -----------------------------------------------------------------------------
ws2 = wb.create_sheet(title='Lieux Sauvages')
ws2.views.sheetView[0].showGridLines = True

headers2 = [
    'Région', 'Lieu', 'Type de Zone', 'N° Dex', 'Pokémon (FR)',
    'Niveau Min', 'Niveau Max', 'Rareté', 'Badge Requis'
]
ws2.append(headers2)
for col_num, _ in enumerate(headers2, 1):
    cell = ws2.cell(row=1, column=col_num)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 28

pokemon_by_id = {p['dex']: p for p in pokemon_list}

for loc in locations_list:
    reg = loc.get('region', 'kanto').capitalize()
    loc_name = loc.get('name', loc['id'])
    loc_type = loc.get('type', 'route').capitalize()
    badge_req = loc.get('badgeReq', 0)
    wilds = loc.get('wild', [])

    if not wilds:
        continue

    for w in wilds:
        pid = int(w[0])
        p_obj = pokemon_by_id.get(pid)
        p_name = p_obj['nameFr'] if p_obj else f"Pokemon #{pid}"
        min_lvl = int(w[1]) if len(w) > 1 else 1
        max_lvl = int(w[2]) if len(w) > 2 else min_lvl
        rarity = w[3] if len(w) > 3 else 'common'
        rarity_fr = 'Commun' if rarity == 'common' else ('Rare' if rarity == 'rare' else 'Peu commun')

        if pid == 249 and loc.get('id') == 'whirlislands':
            loc_name_display = f"{loc_name} (Après quête #139)"
        elif pid == 250 and loc.get('id') == 'tintower':
            loc_name_display = f"{loc_name} (Après quête #140)"
        else:
            loc_name_display = loc_name

        row_data = [
            reg, loc_name_display, loc_type, f"#{pid:03d}", p_name,
            min_lvl, max_lvl, rarity_fr, badge_req
        ]
        ws2.append(row_data)
        cur_r = ws2.max_row
        fill_to_use = ZEBRA_EVEN if cur_r % 2 == 0 else ZEBRA_ODD
        for c_idx in range(1, len(row_data) + 1):
            c = ws2.cell(row=cur_r, column=c_idx)
            c.fill = fill_to_use
            c.font = Font(name=FONT_NAME, size=10)
            c.border = BORDER_THIN
            if c_idx in [1, 3, 4, 6, 7, 8, 9]:
                c.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c.alignment = Alignment(horizontal='left', vertical='center')

# Add roaming legendaries to Sheet 2
roamers_list = [
    ('Kanto', 'Errant / Roaming Kanto (Routes extérieures)', 'Errant', 144, 75, 75, 'Rare (1%)', 0),
    ('Kanto', 'Errant / Roaming Kanto (Routes extérieures)', 'Errant', 145, 75, 75, 'Rare (1%)', 0),
    ('Kanto', 'Errant / Roaming Kanto (Routes extérieures)', 'Errant', 146, 75, 75, 'Rare (1%)', 0),
    ('Kanto', 'Errant / Roaming Kanto (Routes extérieures)', 'Errant', 151, 75, 75, 'Rare (1%)', 0),
    ('Johto', 'Errant / Roaming Johto (Routes extérieures)', 'Errant', 243, 75, 75, 'Rare (1%)', 0),
    ('Johto', 'Errant / Roaming Johto (Routes extérieures)', 'Errant', 244, 75, 75, 'Rare (1%)', 0),
    ('Johto', 'Errant / Roaming Johto (Routes extérieures)', 'Errant', 245, 75, 75, 'Rare (1%)', 0),
    ('Johto', 'Errant / Roaming Johto (Routes extérieures)', 'Errant', 251, 75, 75, 'Rare (1%)', 0),
    ('Hoenn', 'Errant / Roaming Hoenn (Routes extérieures)', 'Errant', 380, 75, 75, 'Rare (1%)', 0),
    ('Hoenn', 'Errant / Roaming Hoenn (Routes extérieures)', 'Errant', 381, 75, 75, 'Rare (1%)', 0),
    ('Hoenn', 'Errant / Roaming Hoenn (Routes extérieures)', 'Errant', 385, 75, 75, 'Rare (1%)', 0),
    ('Hoenn', 'Errant / Roaming Hoenn (Routes extérieures)', 'Errant', 386, 75, 75, 'Rare (1%)', 0),
]

for r_reg, r_loc, r_type, r_pid, r_min, r_max, r_rar, r_badge in roamers_list:
    p_obj = pokemon_by_id.get(r_pid)
    p_name = p_obj['nameFr'] if p_obj else f"Pokemon #{r_pid}"
    row_data = [
        r_reg, r_loc, r_type, f"#{r_pid:03d}", p_name,
        r_min, r_max, r_rar, r_badge
    ]
    ws2.append(row_data)
    cur_r = ws2.max_row
    fill_to_use = ZEBRA_EVEN if cur_r % 2 == 0 else ZEBRA_ODD
    for c_idx in range(1, len(row_data) + 1):
        c = ws2.cell(row=cur_r, column=c_idx)
        c.fill = fill_to_use
        c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER_THIN
        if c_idx in [1, 3, 4, 6, 7, 8, 9]:
            c.alignment = Alignment(horizontal='center', vertical='center')
        else:
            c.alignment = Alignment(horizontal='left', vertical='center')

ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = ws2.dimensions

# -----------------------------------------------------------------------------
# SHEET 3: Évolutions (Evolution Trees)
# -----------------------------------------------------------------------------
ws3 = wb.create_sheet(title='Évolutions')
ws3.views.sheetView[0].showGridLines = True

headers3 = [
    'N° Dex Base', 'Pokémon Base', 'N° Dex Évolution', 'Pokémon Évolution',
    'Méthode d\'Évolution', 'Condition / Niveau / Pierre Requise'
]
ws3.append(headers3)
for col_num, _ in enumerate(headers3, 1):
    cell = ws3.cell(row=1, column=col_num)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 28

for p in pokemon_list:
    if p['evolution']:
        evo = p['evolution']
        from_id = evo['fromId']
        from_p = pokemon_by_id.get(from_id)
        from_name = from_p['nameFr'] if from_p else f"#{from_id}"
        m_type = 'Niveau' if evo['type'] == 'level' else ('Pierre' if evo['type'] == 'stone' else 'Spéciale')
        row_data = [
            f"#{from_id:03d}", from_name, f"#{p['dex']:03d}", p['nameFr'],
            m_type, evo['details']
        ]
        ws3.append(row_data)
        cur_r = ws3.max_row
        fill_to_use = ZEBRA_EVEN if cur_r % 2 == 0 else ZEBRA_ODD
        for c_idx in range(1, len(row_data) + 1):
            c = ws3.cell(row=cur_r, column=c_idx)
            c.fill = fill_to_use
            c.font = Font(name=FONT_NAME, size=10)
            c.border = BORDER_THIN
            if c_idx in [1, 3, 5]:
                c.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c.alignment = Alignment(horizontal='left', vertical='center')

ws3.freeze_panes = 'A2'
ws3.auto_filter.ref = ws3.dimensions

# -----------------------------------------------------------------------------
# SHEET 4: Quêtes & Légendaires
# -----------------------------------------------------------------------------
ws4 = wb.create_sheet(title='Quêtes & Légendaires')
ws4.views.sheetView[0].showGridLines = True

headers4 = [
    'Catégorie', 'ID Quête', 'Titre de la Quête', 'Région',
    'N° Dex Récompense', 'Pokémon Récompensé', 'Niveau Offert', 'Chromatique (Shiny)'
]
ws4.append(headers4)
for col_num, _ in enumerate(headers4, 1):
    cell = ws4.cell(row=1, column=col_num)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
ws4.row_dimensions[1].height = 28

for p in pokemon_list:
    if p['quests']:
        for q in p['quests']:
            cat_fr = 'Histoire' if q['cat'] == 'main' else ('Secondaire' if q['cat'] == 'side' else 'Légendaire')
            row_data = [
                cat_fr, str(q['id']), q['title'], (q['region'] or 'Kanto').capitalize(),
                f"#{p['dex']:03d}", p['nameFr'], q['level'], 'Oui ★' if q['shiny'] else 'Non'
            ]
            ws4.append(row_data)
            cur_r = ws4.max_row
            fill_to_use = ZEBRA_EVEN if cur_r % 2 == 0 else ZEBRA_ODD
            for c_idx in range(1, len(row_data) + 1):
                c = ws4.cell(row=cur_r, column=c_idx)
                c.fill = fill_to_use
                c.font = Font(name=FONT_NAME, size=10)
                c.border = BORDER_THIN
                if c_idx in [1, 2, 4, 5, 7, 8]:
                    c.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    c.alignment = Alignment(horizontal='left', vertical='center')

ws4.freeze_panes = 'A2'
ws4.auto_filter.ref = ws4.dimensions

# -----------------------------------------------------------------------------
# SHEET 5: Marché Noir, Fossiles & Pension
# -----------------------------------------------------------------------------
ws5 = wb.create_sheet(title='Marché Noir & Spéciaux')
ws5.views.sheetView[0].showGridLines = True

headers5 = [
    'Système', 'Région / Source', 'N° Dex', 'Pokémon', 'Condition / Détail / Prix'
]
ws5.append(headers5)
for col_num, _ in enumerate(headers5, 1):
    cell = ws5.cell(row=1, column=col_num)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
ws5.row_dimensions[1].height = 28

# Market rows
for p in pokemon_list:
    if p['market']:
        row_data = [
            'Marché Noir', p['market']['region'].capitalize(),
            f"#{p['dex']:03d}", p['nameFr'], f"Achat direct : {p['market']['price']:,} ₽".replace(',', ' ')
        ]
        ws5.append(row_data)

# Fossil rows
for p in pokemon_list:
    if p['fossil']:
        row_data = [
            'Fossile', 'Mine / Laboratoire',
            f"#{p['dex']:03d}", p['nameFr'], f"Réanimation de l'objet : {p['fossil']['name']}"
        ]
        ws5.append(row_data)

# Breeding rows
for p in pokemon_list:
    if p['breeding']:
        row_data = [
            'Pension (Bébé)', 'Hatchery / Pension',
            f"#{p['dex']:03d}", p['nameFr'], f"Accouplement de : {p['breeding']['parents']}"
        ]
        ws5.append(row_data)

# Starter rows
for p in pokemon_list:
    if p['isStarter']:
        row_data = [
            'Starter', p['region'],
            f"#{p['dex']:03d}", p['nameFr'], "Choix de départ de la région (ou Marché Noir)"
        ]
        ws5.append(row_data)

for cur_r in range(2, ws5.max_row + 1):
    fill_to_use = ZEBRA_EVEN if cur_r % 2 == 0 else ZEBRA_ODD
    for c_idx in range(1, 6):
        c = ws5.cell(row=cur_r, column=c_idx)
        c.fill = fill_to_use
        c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER_THIN
        if c_idx in [1, 2, 3]:
            c.alignment = Alignment(horizontal='center', vertical='center')
        else:
            c.alignment = Alignment(horizontal='left', vertical='center')

ws5.freeze_panes = 'A2'
ws5.auto_filter.ref = ws5.dimensions

# Auto-adjust column widths on all sheets
for sheet in wb.worksheets:
    for col in sheet.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Cap column width between 12 and 60
        sheet.column_dimensions[col_letter].width = max(12, min(max_len + 3, 60))

os.makedirs('docs', exist_ok=True)
output_path = 'docs/POKEDEX_OBTENTION.xlsx'
wb.save(output_path)
print(f"Successfully generated {output_path} with {len(wb.sheetnames)} sheets.")
